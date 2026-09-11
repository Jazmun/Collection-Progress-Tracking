import streamlit as st
import pypdf
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
import re
from datetime import datetime, timedelta

st.set_page_config(page_title="Invoice to Excel Converter", page_icon="📑", layout="wide")

st.title("📑 Landscaping Invoice to Excel Converter")
st.caption("🚀 Version 2.5 — mm/dd/yyyy Date Formatting Active")
st.write("Upload an invoice PDF to extract line items and export directly into your accounting import spreadsheet.")

uploaded_file = st.file_uploader("Choose an Invoice PDF", type=["pdf"])

def extract_invoice_number(text):
    """Detects invoice number across different template styles."""
    # Alternate layout: Date Invoice No. \n 09/10/26 1246 OR Invoice No. 1246
    m2 = re.search(r'Invoice\s*(?:No\.?|#)?\s*(?:\n|\s)+(?:[0-9/]+\s+)?(\d{4,})', text, re.IGNORECASE)
    if m2:
        return m2.group(1)
    # Standard layout: Invoice 1245
    m1 = re.search(r'Invoice\s+(\d{4,})', text, re.IGNORECASE)
    if m1:
        return m1.group(1)
    return None

def clean_description(desc_text):
    """Cleans up the standard description block and removes table headers/footers."""
    lines = desc_text.split('\n')
    cleaned = []
    for l in lines:
        s = l.strip()
        if not s:
            continue
        if "Description Qty / UOM" in s or s == "Description":
            continue
        if "713-657-0875" in s or "lasallelandscaping.com" in s:
            continue
        m_amt = re.search(r'\s+\$([\d,]+\.\d{2})$', s)
        if m_amt:
            s = s[:m_amt.start()].strip()
            if not s:
                continue
        cleaned.append(s)
    return "\n".join(cleaned)

def parse_date(date_str):
    """Parses various date string formats and returns a datetime object."""
    date_str = date_str.strip()
    for fmt in ("%m/%d/%y", "%m/%d/%Y", "%m-%d-%y", "%m-%d-%Y"):
        try:
            return datetime.strptime(date_str, fmt)
        except ValueError:
            pass
    return None

def format_to_mm_dd_yyyy(dt):
    """Formats datetime object to mm/dd/yyyy format with slashes and leading zeroes."""
    if isinstance(dt, datetime):
        return dt.strftime("%m/%d/%Y")
    return str(dt)

def parse_invoices(pdf_bytes):
    reader = pypdf.PdfReader(BytesIO(pdf_bytes))
    invoices = []
    current_inv = None

    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        inv_num = extract_invoice_number(text)
        
        if inv_num:
            if current_inv:
                invoices.append(current_inv)
            current_inv = {
                'inv_num': inv_num,
                'pages': [text]
            }
        else:
            if current_inv:
                current_inv['pages'].append(text)

    if current_inv:
        invoices.append(current_inv)

    records = []

    for inv in invoices:
        full_text = "\n".join(inv['pages'])
        inv_num = inv['inv_num']

        inv_date_str = ""
        due_date_str = ""

        # Format 2 (Alternate layout)
        f2_date = re.search(r'Date\s+Invoice\s+No\.\s*\n\s*([0-9/]+)', full_text, re.IGNORECASE)
        f2_due = re.search(r'Due\s+Date\s*\n\s*.*?\s*([0-9]{1,2}/[0-9]{1,2}/[0-9]{2,4})', full_text, re.IGNORECASE)

        if f2_date:
            dt = parse_date(f2_date.group(1))
            if dt:
                inv_date_str = format_to_mm_dd_yyyy(dt)
            else:
                inv_date_str = f2_date.group(1).strip()

            if f2_due:
                due_dt = parse_date(f2_due.group(1))
                if due_dt:
                    due_date_str = format_to_mm_dd_yyyy(due_dt)
                else:
                    due_date_str = f2_due.group(1).strip()
        else:
            # Format 1 (Standard layout)
            date_match = re.search(r'Date\s+PO#\s*\n\s*([0-9/]+)', full_text)
            if date_match:
                dt = parse_date(date_match.group(1))
                if dt:
                    inv_date_str = format_to_mm_dd_yyyy(dt)
                    if "Due on Receipt" in full_text:
                        due_date_str = inv_date_str
                    else:
                        due_dt = dt + timedelta(days=30)
                        due_date_str = format_to_mm_dd_yyyy(due_dt)
                else:
                    inv_date_str = date_match.group(1).strip()
                    due_date_str = inv_date_str

        # Customer: Extract first line of Property Address
        cust = ""
        if "BILL TO PROPERTY" in full_text:
            bt_idx = full_text.find("BILL TO PROPERTY")
            end_idx = full_text.find("Amount Due", bt_idx)
            if end_idx == -1:
                end_idx = full_text.find("Please detach", bt_idx)
            block = full_text[bt_idx + len("BILL TO PROPERTY"):end_idx]
            lines = [l.strip() for l in block.split('\n') if l.strip()]
            found_zip = -1
            for i, l in enumerate(lines):
                if re.search(r'[A-Z]{2}\s+\d{5}', l):
                    found_zip = i
                    break
            if found_zip != -1 and found_zip + 1 < len(lines):
                cust = lines[found_zip + 1]
            elif lines:
                cust = lines[0]
        else:
            bt_idx = full_text.find("Bill To Property Address")
            desc_idx = full_text.find("Description", bt_idx) if bt_idx != -1 else -1
            if bt_idx != -1 and desc_idx != -1:
                addr_block = full_text[bt_idx + len("Bill To Property Address"):desc_idx]
                addr_lines = [l.strip() for l in addr_block.split('\n') if l.strip()]
                found_zip_idx = -1
                for i, l in enumerate(addr_lines):
                    if re.search(r'[A-Z]{2}\s+\d{5}', l):
                        found_zip_idx = i
                        break
                if found_zip_idx != -1 and found_zip_idx + 1 < len(addr_lines):
                    cust = addr_lines[found_zip_idx + 1]
                elif addr_lines:
                    cust = addr_lines[0]

        # Unit Price & Tax
        unit_price = 0.0
        has_tax = "No"

        if "BILL TO PROPERTY" in full_text:
            tot_match = re.search(r'Total\s*\$?([\d,]+\.\d{2})\s+\$?([\d,]+\.\d{2})', full_text)
            if tot_match:
                unit_price = float(tot_match.group(1).replace(',', ''))
                tax_amt = float(tot_match.group(2).replace(',', ''))
                if tax_amt > 0:
                    has_tax = "Yes"
            else:
                amt_match = re.search(r'EXT PRICE\s*\n?\s*\$?([\d,]+\.\d{2})', full_text)
                if amt_match:
                    unit_price = float(amt_match.group(1).replace(',', ''))
        else:
            sub_match = re.search(r'Subtotal\s*\$?([\d,]+\.\d{2})', full_text)
            unit_price = float(sub_match.group(1).replace(',', '')) if sub_match else 0.0

            tax_match = re.search(r'Sales Tax\s*\$?([\d,]+\.\d{2})', full_text)
            if tax_match:
                t_val = float(tax_match.group(1).replace(',', ''))
                if t_val > 0:
                    has_tax = "Yes"

        # Description
        clean_desc = ""
        if "BILL TO PROPERTY" in full_text:
            start_desc = full_text.find("Please detach top portion and return with your payment.")
            if start_desc != -1:
                start_desc += len("Please detach top portion and return with your payment.")
            else:
                start_desc = full_text.find("QTY ITEM")
                if start_desc != -1:
                    start_desc += len("QTY ITEM")

            end_desc = full_text.find("Total", start_desc)
            if end_desc == -1:
                end_desc = full_text.find("UNIT PRICE", start_desc)

            raw_desc = full_text[start_desc:end_desc] if end_desc != -1 else full_text[start_desc:]
            lines = raw_desc.split('\n')
            c_lines = []
            for l in lines:
                s = l.strip()
                if not s or "QTY ITEM" in s or "UNIT PRICE" in s or "EXT PRICE" in s:
                    continue
                if re.match(r'^\$?[\d,]+\.\d{2}(\s+\$?[\d,]+\.\d{2})*$', s):
                    continue
                s = re.sub(r'\s+\$?[\d,]+\.\d{2}.*$', '', s).strip()
                if s:
                    c_lines.append(s)
            clean_desc = "\n".join(c_lines)
        else:
            desc_start = full_text.find("Description Qty / UOM")
            if desc_start == -1:
                desc_start = full_text.find("Description")
            sub_start = full_text.find("Subtotal", desc_start)
            raw_desc = full_text[desc_start:sub_start] if sub_start != -1 else full_text[desc_start:]
            clean_desc = clean_description(raw_desc)

        records.append({
            "Post": "Yes",
            "Invoice Date": inv_date_str,
            "Due Date": due_date_str,
            "Invoice Number": inv_num,
            "Transaction Type": "Invoice",
            "Customer": cust,
            "Vendor": "",
            "Currency Code": "",
            "Products/Services": "Side Jobs",
            "Description": clean_desc,
            "Qty": 1,
            "Discount %": "",
            "Unit Price": unit_price,
            "Category": "",
            "Location": "",
            "Class": "",
            "Tax": has_tax
        })

    return records

def create_excel(records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice Import"

    headers = [
        "Post", "Invoice Date", "Due Date", "Invoice Number", "Transaction Type",
        "Customer", "Vendor", "Currency Code", "Products/Services", "Description",
        "Qty", "Discount %", "Unit Price", "Category", "Location", "Class", "Tax"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, r in enumerate(records, start=2):
        row_values = [r[h] for h in headers]
        ws.append(row_values)
        fill_color = "F9FAFC" if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        for col_idx, cell in enumerate(ws[row_idx], start=1):
            cell.fill = row_fill
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            col_name = headers[col_idx - 1]
            if col_name in ["Post", "Transaction Type", "Qty", "Tax", "Invoice Date", "Due Date", "Invoice Number"]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_name == "Unit Price":
                cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="top")
            elif col_name == "Description":
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '').split('\n')[0]) for cell in col)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 42), 12)
    ws.column_dimensions['J'].width = 50

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

if uploaded_file is not None:
    with st.spinner("Processing PDF and extracting line items..."):
        data = parse_invoices(uploaded_file.read())

    if data:
        st.success(f"Successfully processed {len(data)} invoices!")
        
        # Display preview table with dates, customer, and amounts
        preview_data = [{
            "Invoice #": r["Invoice Number"], 
            "Invoice Date": r["Invoice Date"],
            "Due Date": r["Due Date"],
            "Customer": r["Customer"], 
            "Amount": f"${r['Unit Price']:,.2f}"
        } for r in data]
        st.table(preview_data)
        
        excel_data = create_excel(data)

        st.download_button(
            label="📥 Download Excel Spreadsheet",
            data=excel_data,
            file_name="Extracted_Invoices.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    else:
        st.error("No invoice data found in the uploaded file.")
